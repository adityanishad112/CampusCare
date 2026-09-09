import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors (CampusCare Palette)
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Primary Header
    BLUE = RGBColor(59, 130, 246)      # #3B82F6 - Primary Accent
    DARK = RGBColor(15, 23, 42)        # #0F172A - Body text
    MUTED = RGBColor(100, 116, 139)    # #64748B - Subtitle / Secondary
    BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC - Card background
    WHITE = RGBColor(255, 255, 255)
    TEAL = RGBColor(13, 148, 136)      # #0D9488 - Success/Resolved
    AMBER = RGBColor(217, 119, 6)      # #D97706 - Warning
    CARD_BORDER = RGBColor(226, 232, 240)

    def add_header(slide, title_text, category_text="CAMPUSCARE — CSE CAPSTONE PROJECT"):
        # Category / Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=BG_LIGHT, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark / Navy Theme)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Title Card
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FOURTH YEAR CSE CAPSTONE PROJECT"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(147, 197, 253) # Light blue
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "CampusCare — AI-Based Campus Complaint Management System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "An intelligent, auditable grievance resolution ecosystem featuring real-time NLP department classification, heuristic safety triage, and vector similarity duplicate suppression."
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_after = Pt(36)

    p3 = tf1.add_paragraph()
    p3.text = "Tech Stack: React 19 • TypeScript • Tailwind CSS • FastAPI • scikit-learn • PostgreSQL"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(96, 165, 250)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Motivation
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Problem Statement & Operational Challenges")

    # 3 Column Cards
    col_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    top_y = Inches(1.8)
    card_h = Inches(4.8)

    problems = [
        ("Manual Triage Bottlenecks",
         "Grievances are logged via manual registers, paper forms, or generic emails.",
         [
             "Complaints sit in incorrect department inboxes for days.",
             "High misclassification rate during manual assignment.",
             "Technicians lack timely notifications to initiate repairs."
         ],
         AMBER),
        ("Duplicate Ticket Flooding",
         "Localized outages trigger an influx of redundant, uncoordinated submissions.",
         [
             "A single lab Wi-Fi outage triggers dozens of duplicate tickets.",
             "Staff spend hours filtering repeat reports instead of fixing root causes.",
             "Distorts university maintenance analytics and backlog figures."
         ],
         RGBColor(225, 29, 72)),
        ("Zero Auditability & Feedback",
         "Students and campus leadership operate with zero transparent visibility.",
         [
             "Students have no progress timeline or assigned technician info.",
             "No structured way to dispute superficial or incomplete closures.",
             "Administration lacks SLA metrics, resolution velocity, and KPI dashboards."
         ],
         NAVY)
    ]

    for i, (title, subtitle, bullets, accent) in enumerate(problems):
        x = start_x + i * (col_w + gap)
        add_card(slide2, x, top_y, col_w, card_h)

        # Accent top bar
        bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, col_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tbox = slide2.shapes.add_textbox(x + Inches(0.3), top_y + Inches(0.3), col_w - Inches(0.6), card_h - Inches(0.5))
        tf = tbox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(8)

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = MUTED
        p_sub.space_after = Pt(16)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(13)
            pb.font.color.rgb = DARK
            pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution & Key Objectives
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Proposed Solution: The CampusCare Ecosystem")

    # Left Column: Overview card
    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tbox_l = slide3.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_l = tbox_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "Core Design Philosophy"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    points_l = [
        ("Automated Intelligence First", "Employ machine learning to assist human operators instantly at the point of submission, eliminating triage queues."),
        ("Strict Deterministic Lifecycle", "Enforce an auditable state machine with immutable logs—preventing undocumented closures or unauthorized status skips."),
        ("Human-in-the-Loop Governance", "Retain administrative oversight with automated queues for borderline or low-confidence ML classifications."),
        ("Accessible Modern UX", "Responsive university design system built to WCAG 2.1 standards with 1-click evaluator test personas.")
    ]

    for title, desc in points_l:
        p_t = tf_l.add_paragraph()
        p_t.text = "✔ " + title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = BLUE

        p_d = tf_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK
        p_d.space_after = Pt(10)

    # Right Column: 3 Objectives Cards
    right_x = Inches(6.8)
    obj_cards = [
        ("1. Intelligent Automated Triage", "Classify unstructured complaint text into 7 university departments using TF-IDF and Logistic Regression with >94% F1-score.", TEAL),
        ("2. Real-Time Duplicate Prevention", "Compute vector cosine similarity over existing open complaints to warn students before creating duplicate tickets.", BLUE),
        ("3. End-to-End Operational Transparency", "Provide visual audit timelines, 5-star student satisfaction ratings, 7-day reopening grace periods, and executive CSV exports.", NAVY)
    ]

    for i, (title, desc, color) in enumerate(obj_cards):
        y = Inches(1.8) + i * Inches(1.65)
        add_card(slide3, right_x, y, Inches(5.7), Inches(1.45))
        bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, y, Inches(0.12), Inches(1.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = slide3.shapes.add_textbox(right_x + Inches(0.3), y + Inches(0.2), Inches(5.2), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 4: System Architecture & Technical Stack
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Three-Tier Decoupled System Architecture")

    # 3 Layers horizontally
    layer_w = Inches(3.64)
    start_x = Inches(0.8)
    layer_h = Inches(4.8)

    layers = [
        ("Presentation Layer", "Client Web Application", [
            "React 19 & TypeScript: Type-safe, component-driven UI architecture.",
            "Tailwind CSS v3.4: Atomic design system tokens (Deep Navy #1E40AF & Service Blue #3B82F6).",
            "Lucide React & Recharts: Visual analytics and iconography.",
            "WCAG 2.1 Compliant: High contrast typography (Atkinson Hyperlegible)."
        ], BLUE),
        ("Application & Intelligence", "FastAPI Core Services", [
            "FastAPI Python: Async RESTful microservice with sub-5ms latency.",
            "ML Classifier: TF-IDF (1,2-grams) + Multinomial Logistic Regression.",
            "Safety Priority Engine: Deterministic keyword heuristics.",
            "Duplicate Engine: Vector cosine similarity thresholding.",
            "Security: Stateless HMAC-SHA256 JWT & bcrypt password hashing."
        ], NAVY),
        ("Data Persistence Layer", "Relational Storage & Schema", [
            "SQLAlchemy ORM: Type-checked database querying and relationships.",
            "Alembic Migrations: Version-controlled database schema evolution.",
            "PostgreSQL 16: Robust, ACID-compliant production database in Docker.",
            "SQLite Fallback: Zero-config local development runtime without external setup."
        ], TEAL)
    ]

    for i, (layer_name, subtitle, bullets, color) in enumerate(layers):
        x = start_x + i * (layer_w + gap)
        add_card(slide4, x, top_y, layer_w, layer_h)

        bar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, layer_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = slide4.shapes.add_textbox(x + Inches(0.3), top_y + Inches(0.3), layer_w - Inches(0.6), layer_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = layer_name
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(4)

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = MUTED
        p_sub.space_after = Pt(16)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(12.5)
            pb.font.color.rgb = DARK
            pb.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 5: Finite State Machine & Ticket Lifecycle
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Finite State Machine & Workflow Engine")

    # Workflow Steps
    steps = [
        ("1. Submitted", "Complaint filed by student. Automated ML category & priority assigned.", BLUE),
        ("2. Assigned", "Staff claims ticket or Admin assigns to department technician.", RGBColor(99, 102, 241)),
        ("3. In Progress", "Staff initiates inspection/work. Internal notes recorded.", AMBER),
        ("4. Resolved", "Staff marks work complete with mandatory resolution documentation.", TEAL),
        ("5. Closed", "Student verifies fix, submits 1-5 star rating & feedback.", RGBColor(22, 163, 74)),
        ("6. Reopened", "Student rejects resolution within 7 days. Returns to active queue.", RGBColor(225, 29, 72))
    ]

    step_w = Inches(3.64)
    step_h = Inches(2.2)

    for i, (stitle, sdesc, scolor) in enumerate(steps):
        row = i // 3
        col = i % 3
        x = start_x + col * (step_w + gap)
        y = top_y + row * (step_h + Inches(0.35))

        add_card(slide5, x, y, step_w, step_h)
        bar = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, step_w, Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = scolor
        bar.line.fill.background()

        tb = slide5.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), step_w - Inches(0.5), step_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 6: Machine Learning Pipeline & Mathematics
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Applied Machine Learning Classification Pipeline")

    # Left Column: Steps
    add_card(slide6, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.8))
    tb_l6 = slide6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.4), Inches(4.3))
    tf_l6 = tb_l6.text_frame
    tf_l6.word_wrap = True

    p = tf_l6.paragraphs[0]
    p.text = "Natural Language Processing Methodology"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    ml_steps = [
        ("Text Normalization", "Case folding, alphanumeric filtering, and domain-specific campus stopword cleaning."),
        ("Sublinear TF-IDF Vectorization", "Extracts unigrams and bigrams (1,2) with logarithmic term-frequency scaling to prevent lengthy descriptions from skewing short titles."),
        ("Multinomial Logistic Regression", "Trains with L2 regularization (C=1.0) outputting calibrated posterior class probabilities across 7 departments."),
        ("Confidence Calibration & Thresholding", "Returns top predicted category with confidence score P(y=k|x). Tickets with confidence < 60% are flagged for Human-in-the-Loop review.")
    ]

    for title, desc in ml_steps:
        pt = tf_l6.add_paragraph()
        pt.text = "• " + title
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = BLUE

        pd = tf_l6.add_paragraph()
        pd.text = "   " + desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = DARK
        pd.space_after = Pt(8)

    # Right Column: 7 Categories
    add_card(slide6, Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.8))
    tb_r6 = slide6.shapes.add_textbox(Inches(7.4), Inches(2.0), Inches(4.8), Inches(4.3))
    tf_r6 = tb_r6.text_frame
    tf_r6.word_wrap = True

    p = tf_r6.paragraphs[0]
    p.text = "7 Institutional Target Categories"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    cats = [
        ("Hostel", "Plumbing, hot water geysers, cleanliness, furniture, corridor noise."),
        ("Mess / Canteen", "Food quality, dining hygiene, cutlery, menu pricing."),
        ("IT & Network", "Wi-Fi coverage, portal login, lab switchboards, server latency."),
        ("Electrical", "Switchboard sparks, power cuts, tube lights, elevator stoppage."),
        ("Academic", "Classroom projectors, microphones, syllabus schedules."),
        ("Infrastructure", "Potholes, broken benches, water purifiers, window latches."),
        ("Security", "Unauthorized campus entry, parking blockage, perimeter lights.")
    ]

    for cname, cdesc in cats:
        p_c = tf_r6.add_paragraph()
        p_c.text = "▸ " + cname + ": "
        p_c.font.size = Pt(12)
        p_c.font.bold = True
        p_c.font.color.rgb = NAVY
        p_c.space_after = Pt(4)

        # Add description inline
        run = p_c.add_run()
        run.text = cdesc
        run.font.bold = False
        run.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 7: Experimental Evaluation & Quantitative Results
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Experimental Evaluation & Quantitative Results")

    # 4 Metric KPI Cards
    kpis = [
        ("94.89%", "Macro F1-Score", TEAL),
        ("96.10%", "Macro Precision", BLUE),
        ("95.24%", "Macro Recall", NAVY),
        ("< 5 ms", "Inference Latency", RGBColor(16, 185, 129))
    ]

    kw = Inches(2.65)
    kgap = Inches(0.35)
    for i, (val, label, col) in enumerate(kpis):
        kx = Inches(0.8) + i * (kw + kgap)
        add_card(slide7, kx, Inches(1.8), kw, Inches(1.3))

        tb = slide7.shapes.add_textbox(kx, Inches(1.9), kw, Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = val
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = label
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED

    # Table of Category Results
    t_top = Inches(3.4)
    t_left = Inches(0.8)
    t_w = Inches(11.7)
    t_h = Inches(3.2)

    rows, cols = 8, 4
    table_shape = slide7.shapes.add_table(rows, cols, t_left, t_top, t_w, t_h)
    table = table_shape.table

    table.columns[0].width = Inches(4.5)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(2.4)
    table.columns[3].width = Inches(2.4)

    headers = ["Department Category", "Precision", "Recall", "F1-Score"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.text = h
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("Hostel Maintenance", "0.92", "1.00", "0.96"),
        ("IT & Network Infrastructure", "1.00", "0.94", "0.97"),
        ("Electrical Maintenance", "0.95", "1.00", "0.97"),
        ("Mess & Canteen Sanitation", "1.00", "0.89", "0.94"),
        ("Academic & Classroom Facilities", "1.00", "0.88", "0.93"),
        ("General Infrastructure", "0.91", "1.00", "0.95"),
        ("Campus Security & Safety", "0.95", "0.95", "0.95"),
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)
            p.font.color.rgb = DARK
            if j > 0:
                p.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 8: Priority Detection & Safety Hazards
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "Heuristic Priority & Emergency Hazard Detection")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_l8 = slide8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_l8 = tb_l8.text_frame
    tf_l8.word_wrap = True

    p = tf_l8.paragraphs[0]
    p.text = "The Safety Guarantee Principle"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    p_body = tf_l8.add_paragraph()
    p_body.text = "Purely statistical models can fail on rare but life-threatening campus hazards due to training distribution imbalance.\n\nCampusCare couples the statistical ML classifier with a deterministic Priority Rule Engine that scans for critical physical and infrastructural triggers."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = DARK
    p_body.space_after = Pt(16)

    p_sub = tf_l8.add_paragraph()
    p_sub.text = "Key Benefits:"
    p_sub.font.size = Pt(14)
    p_sub.font.bold = True
    p_sub.font.color.rgb = BLUE

    b_pts = [
        "Instant Critical elevation for fires, gas leaks, and electric sparks.",
        "Generates explainable rationale displayed directly to students and staff.",
        "Enables instant visual badges and priority filtering on staff queues."
    ]
    for b in b_pts:
        pb = tf_l8.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(12)
        pb.font.color.rgb = DARK
        pb.space_after = Pt(6)

    # Right: Priority Levels Breakdown
    right_x = Inches(6.8)
    levels = [
        ("Critical", "Sparks, fire, smoke, gas leak, building collapse risk, severe injury hazard.", RGBColor(225, 29, 72)),
        ("High", "Elevator stoppage, total hostel water cut, campus-wide Wi-Fi down during exams.", RGBColor(234, 88, 12)),
        ("Medium", "Single lab workstation issue, broken chair, dripping tap, slow internet.", BLUE),
        ("Low", "Aesthetic paint peel, suggestions, non-urgent minor fixture alignment.", RGBColor(100, 116, 139))
    ]

    for i, (lvl, desc, col) in enumerate(levels):
        y = Inches(1.8) + i * Inches(1.22)
        add_card(slide8, right_x, y, Inches(5.7), Inches(1.1))
        bar = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, y, Inches(0.12), Inches(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = slide8.shapes.add_textbox(right_x + Inches(0.3), y + Inches(0.15), Inches(5.2), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = lvl + " Priority"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 9: Vector Similarity Duplicate Detection
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "Vector Similarity & Duplicate Complaint Suppression")

    # 3 Workflow Blocks
    w_card = Inches(3.64)
    for i, (stage, desc, details) in enumerate([
        ("Stage 1: Context Filter", "Filters candidate tickets by status & location.", [
            "Selects open complaints (Submitted, Assigned, In Progress).",
            "Filters candidates within matching location tokens (e.g. 'Lab 3', 'Hostel B').",
            "Eliminates unnecessary global comparisons to maximize speed."
        ]),
        ("Stage 2: TF-IDF Embedding", "Transforms text into normalized vector space.", [
            "Converts candidate title & description into sparse TF-IDF vectors.",
            "Normalizes vectors using L2 norm to ensure invariant length scaling.",
            "Preserves domain vocabulary weights for campus facilities."
        ]),
        ("Stage 3: Cosine Similarity", "Evaluates cosine angle against threshold.", [
            "Computes dot product: Similarity = (u · v) / (||u|| ||v||).",
            "Threshold T >= 0.65 triggers real-time pre-submission warning.",
            "Allows staff to link duplicate children to a master parent ticket."
        ])
    ]):
        x = start_x + i * (w_card + gap)
        add_card(slide9, x, top_y, w_card, Inches(4.8))

        bar = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, w_card, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE
        bar.line.fill.background()

        tb = slide9.shapes.add_textbox(x + Inches(0.3), top_y + Inches(0.3), w_card - Inches(0.6), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stage
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(6)

        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = MUTED
        p_sub.space_after = Pt(14)

        for d in details:
            pd = tf.add_paragraph()
            pd.text = "• " + d
            pd.font.size = Pt(12.5)
            pd.font.color.rgb = DARK
            pd.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 10: Security, RBAC & Tenant Data Isolation
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "Security Architecture & Role-Based Access Control")

    # 3 Role Cards
    roles = [
        ("Student Role", "Grievance Filing & Feedback", [
            "Self-registration restricted strictly to student scope.",
            "Can submit grievances, view own complaints, and add public comments.",
            "Rate resolved tickets (1-5 stars) or reopen within 7 days.",
            "Cryptographically restricted from viewing other students' tickets."
        ], BLUE),
        ("Department Staff", "Queue Triage & Resolution", [
            "Strict tenant isolation: view only complaints in assigned department.",
            "Claim unassigned complaints and execute valid status transitions.",
            "Post internal private notes hidden from students.",
            "Mandatory resolution remarks before resolving any ticket."
        ], TEAL),
        ("Administrator", "Global Governance & Analytics", [
            "System-wide visibility across all departments and complaint queues.",
            "Dedicated ML Review Queue for confidence scores < 60%.",
            "Department CRUD and staff user provisioning.",
            "Full CSV audit stream export for campus records."
        ], NAVY)
    ]

    for i, (rtitle, rsub, rbullets, rcol) in enumerate(roles):
        x = start_x + i * (w_card + gap)
        add_card(slide10, x, top_y, w_card, Inches(4.8))

        bar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, w_card, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rcol
        bar.line.fill.background()

        tb = slide10.shapes.add_textbox(x + Inches(0.3), top_y + Inches(0.3), w_card - Inches(0.6), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = rtitle
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(4)

        p_sub = tf.add_paragraph()
        p_sub.text = rsub
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = MUTED
        p_sub.space_after = Pt(14)

        for b in rbullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(12.5)
            pb.font.color.rgb = DARK
            pb.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 11: Student Experience & Triage Flow
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide11, "Student Experience: Smart Grievance Lifecycle")

    features_s = [
        ("Real-Time AI Suggestion", "As the student types title & description, category is predicted dynamically with live confidence gauge.", BLUE),
        ("Duplicate Warning Banner", "Alerts student if an active ticket in the same location already covers this issue, preventing redundant filing.", AMBER),
        ("Visual Status Timeline", "Step-by-step graphical progress tracker showing exact timestamped history from submission to resolution.", TEAL),
        ("Satisfaction Rating & Reopen", "Submit 1-5 star quality feedback, or reopen unresolved issues within 7 days with mandatory rationale.", RGBColor(22, 163, 74))
    ]

    card_w11 = Inches(5.6)
    card_h11 = Inches(2.2)

    for i, (title, desc, col) in enumerate(features_s):
        row = i // 2
        col_idx = i % 2
        x = Inches(0.8) + col_idx * (card_w11 + Inches(0.5))
        y = Inches(1.8) + row * (card_h11 + Inches(0.35))

        add_card(slide11, x, y, card_w11, card_h11)
        bar = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), card_h11)
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = slide11.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), card_w11 - Inches(0.5), card_h11 - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 12: Staff Operations & SLA Management
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide12, "Staff Operations: Queue Triage & Resolution Workflow")

    features_st = [
        ("Department Queue Isolation", "Staff only see complaints assigned to their department. Zero data spillover across department boundaries.", NAVY),
        ("1-Click Ticket Claiming", "Staff members assign tickets to themselves with a single click, instantly moving status to Assigned.", BLUE),
        ("Internal Private Notes", "Private communication channel visible only to department technicians and admins for coordination.", AMBER),
        ("Mandatory Resolution Notes", "System enforces entering technical resolution summary before permitting transition to Resolved.", TEAL)
    ]

    for i, (title, desc, col) in enumerate(features_st):
        row = i // 2
        col_idx = i % 2
        x = Inches(0.8) + col_idx * (card_w11 + Inches(0.5))
        y = Inches(1.8) + row * (card_h11 + Inches(0.35))

        add_card(slide12, x, y, card_w11, card_h11)
        bar = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), card_h11)
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = slide12.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), card_w11 - Inches(0.5), card_h11 - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 13: Administrator Oversight & Review Queue
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide13, "Administrator Oversight & Human-in-the-Loop Governance")

    # Left: ML Review Queue
    add_card(slide13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_l13 = slide13.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_l13 = tb_l13.text_frame
    tf_l13.word_wrap = True

    p = tf_l13.paragraphs[0]
    p.text = "Human-in-the-Loop ML Review Queue"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)

    p_sub = tf_l13.add_paragraph()
    p_sub.text = "Why it matters in real-world engineering:"
    p_sub.font.size = Pt(13)
    p_sub.font.bold = True
    p_sub.font.color.rgb = BLUE
    p_sub.space_after = Pt(8)

    review_bullets = [
        "Vague, colloquial, or ambiguous complaint descriptions naturally yield lower model confidence.",
        "Rather than forcing a low-confidence guess, the system flags tickets with AI confidence < 60%.",
        "Administrators access a dedicated Review Queue to manually inspect and reassign departments with 1 click.",
        "Eliminates misrouting ping-pong between department staff."
    ]
    for b in review_bullets:
        pb = tf_l13.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(12)
        pb.font.color.rgb = DARK
        pb.space_after = Pt(6)

    # Right: Analytics & Reporting
    add_card(slide13, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb_r13 = slide13.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.2))
    tf_r13 = tb_r13.text_frame
    tf_r13.word_wrap = True

    p = tf_r13.paragraphs[0]
    p.text = "Operational Analytics & Governance"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)

    p_sub2 = tf_r13.add_paragraph()
    p_sub2.text = "Key Administrative Capabilities:"
    p_sub2.font.size = Pt(13)
    p_sub2.font.bold = True
    p_sub2.font.color.rgb = TEAL
    p_sub2.space_after = Pt(8)

    admin_bullets = [
        "Real-Time KPIs: Total volume, pending counts, resolution rate %, and mean resolution duration in hours.",
        "Interactive Visualizations: Department workload bar charts, category distribution pie charts.",
        "Department & Staff Provisioning: Create new campus departments and configure staff accounts dynamically.",
        "One-Click CSV Audit Stream: Filter complaints by status or date and download standardized CSV reports for university senate reviews."
    ]
    for b in admin_bullets:
        pb = tf_r13.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(12)
        pb.font.color.rgb = DARK
        pb.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 14: Limitations & Future Enhancements
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide14, "Engineering Limitations & Future Enhancements")

    # 2 Big Columns: Current Limitations vs Future Roadmap
    add_card(slide14, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    bar = slide14.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = AMBER
    bar.line.fill.background()

    tb_l14 = slide14.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_l14 = tb_l14.text_frame
    tf_l14.word_wrap = True

    p = tf_l14.paragraphs[0]
    p.text = "Current System Boundaries"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    limits = [
        ("Text-Only Classification", "Models evaluate textual complaints; uploaded photo attachments are stored but not yet parsed via computer vision."),
        ("English Language Corpus", "Trained primarily on English campus terminology; regional vernacular terms require corpus expansion."),
        ("Periodic Model Updates", "Model retraining is currently executed via batch pipeline scripts rather than real-time continuous streaming.")
    ]
    for t, d in limits:
        pt = tf_l14.add_paragraph()
        pt.text = "⚠ " + t
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = AMBER

        pd = tf_l14.add_paragraph()
        pd.text = d
        pd.font.size = Pt(12)
        pd.font.color.rgb = DARK
        pd.space_after = Pt(8)

    # Right: Roadmap
    add_card(slide14, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    bar = slide14.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    tb_r14 = slide14.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.2))
    tf_r14 = tb_r14.text_frame
    tf_r14.word_wrap = True

    p = tf_r14.paragraphs[0]
    p.text = "Future Research & Engineering Roadmap"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    future = [
        ("Multimodal Vision Verification", "Integrate CNN / Vision-Language models to inspect photos of water damage, broken switches, or cracks."),
        ("Interactive Geospatial Heatmaps", "Render campus floor plans showing live hotspots of recurring failures across blocks and wings."),
        ("Automated WhatsApp Webhooks", "Dispatch instant WhatsApp / SMS alerts to duty technicians when Critical tickets are flagged."),
        ("Active Learning Loop", "Retrain NLP models automatically as administrators verify or reclassify tickets over time.")
    ]
    for t, d in future:
        pt = tf_r14.add_paragraph()
        pt.text = "✦ " + t
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEAL

        pd = tf_r14.add_paragraph()
        pd.text = d
        pd.font.size = Pt(12)
        pd.font.color.rgb = DARK
        pd.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 15: Conclusion & Q&A
    # -------------------------------------------------------------
    slide15 = prs.slides.add_slide(blank_slide_layout)
    bg15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = NAVY
    bg15.line.fill.background()

    tb15 = slide15.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p = tf15.paragraphs[0]
    p.text = "CONCLUSION & SUMMARY"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p.space_after = Pt(12)

    p1 = tf15.add_paragraph()
    p1.text = "CampusCare delivers a production-ready, auditable campus grievance ecosystem."
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(20)

    p2 = tf15.add_paragraph()
    p2.text = "Key Achievements:\n• High-performance ML categorization with 94.89% Macro F1-score and sub-5ms latency\n• Automated hazard priority triage and vector duplicate suppression\n• Strict auditable finite state machine with RBAC and private staff communication\n• Production-ready Docker containerization and zero-config local development"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_after = Pt(30)

    p3 = tf15.add_paragraph()
    p3.text = "Thank you! We welcome your questions."
    p3.font.size = Pt(22)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(96, 165, 250)

    output_path = os.path.join(os.path.dirname(__file__), "CampusCare_Final_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
